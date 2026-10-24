# Workflow Builder runtime: the generic part (no PowerFactory calls).
#
# The Builder copies the helpers a workflow needs out of these runtime files into the one Python
# file it generates. Each helper is a section:
#     # @helper <name>          what blocks and the generator ask for
#     # @needs <a>, <b>         helpers this one calls (copied in before it)
#     # @imports <import line>  one line per import it needs (deduplicated across helpers)
#     ...code...
#     # @end
# Sections are copied in file order, so a helper's code may refer to names defined later only
# inside function bodies. The whole set of runtime files also works concatenated, which is how
# the tests load it. Settings such as OUTPUT_DIR and WORKFLOW_NAME are module globals written
# by the Start block; every helper reads them with a default so it also runs without them.

# @helper log
# @needs run_dir
# @imports import sys
# @imports import time
_LOG = {"fh": None, "t0": time.time(), "warnings": 0, "errors": 0}


def log(msg, level="info"):
    """Print a timestamped line and append it to run.log in the run folder."""
    stamp = time.strftime("%H:%M:%S")
    tag = {"warn": "WARNING ", "error": "ERROR "}.get(level, "")
    if level == "warn":
        _LOG["warnings"] += 1
    elif level == "error":
        _LOG["errors"] += 1
    line = f"[{stamp}] {tag}{msg}"
    if level != "debug":
        try:
            print(line, flush=True)
        except UnicodeEncodeError:
            # A console or pipe that cannot show a name (cp1252 and a Greek letter, say) must not
            # stop the run: show what it can.
            enc = getattr(sys.stdout, "encoding", None) or "ascii"
            print(line.encode(enc, "replace").decode(enc, "replace"), flush=True)
    try:
        if _LOG["fh"] is None:
            _LOG["fh"] = open(os.path.join(run_dir(), "run.log"), "a", encoding="utf-8")
        _LOG["fh"].write(line + "\n")
        _LOG["fh"].flush()
    except Exception:
        pass
# @end

# @helper StopWorkflow
class StopWorkflow(Exception):
    """Raised by a Stop the workflow block; main() catches it and finishes cleanly."""
# @end

# @helper run_dir
# @imports import os
# @imports import re
# @imports import sys
# @imports import time
_RUN = {"dir": None, "files": []}


def safe_filename(text, maxlen=80):
    """A file-name-safe version of text (Windows rules), never empty."""
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", str(text)).strip(" ._")
    s = re.sub(r"\s+", "_", s)
    return (s[:maxlen].rstrip(" ._") or "output")


def script_dir():
    """The folder of the running script (the working folder when that cannot be told)."""
    here = globals().get("__file__") or (sys.argv[0] if sys.argv and sys.argv[0] else "")
    return os.path.dirname(os.path.abspath(here)) if here else os.getcwd()


def run_dir():
    """The folder this run writes to: OUTPUT_DIR/<workflow>_<date>_<time>, made on first use.
    With OUTPUT_DIR empty it goes next to the script, whatever folder the script was started from."""
    if _RUN["dir"] is None:
        base = globals().get("OUTPUT_DIR") or script_dir()
        name = safe_filename(globals().get("WORKFLOW_NAME") or "workflow", 50)
        path = os.path.join(os.path.abspath(os.path.expanduser(base)), name + "_" + time.strftime("%Y%m%d_%H%M%S"))
        n = 1
        while os.path.exists(path + ("" if n == 1 else f"_{n}")):
            n += 1
        path = path + ("" if n == 1 else f"_{n}")
        os.makedirs(path, exist_ok=True)
        _RUN["dir"] = path
    return _RUN["dir"]


def out_file(stem, ext, sub=None):
    """A new file path in the run folder (or its sub folder); the loop labels are added to the
    name when called inside loops, and a counter keeps names unique."""
    folder = run_dir() if not sub else os.path.join(run_dir(), sub)
    os.makedirs(folder, exist_ok=True)
    tag = ""
    ctx_fn = globals().get("ctx_tag")
    if callable(ctx_fn):
        tag = ctx_fn()
    base = safe_filename(stem + ("_" + tag if tag else ""), 120)
    ext = ext.lstrip(".")
    path = os.path.join(folder, f"{base}.{ext}")
    n = 2
    while os.path.exists(path) or path in _RUN["files"]:
        path = os.path.join(folder, f"{base}_{n}.{ext}")
        n += 1
    _RUN["files"].append(path)
    return path
# @end

# @helper ctx
# @needs log
# @imports import time
_CTX = []
_PASSES = {"run": 0, "skipped": 0, "failed": 0, "reasons": [], "durations": {}}


def ctx_push(label, value, i=None, n=None):
    """Start a loop pass: its label and value go into table rows, file names and messages.
    i is the pass number (1-based) and n the number of passes, when known, for the progress line."""
    frame = {"label": str(label), "value": value, "t0": time.time(), "reason": None, "tables": set()}
    _CTX.append(frame)
    _PASSES["run"] += 1
    if i is not None:
        done = _PASSES["durations"].get(frame["label"], [])
        left = ""
        if done and n:
            remaining = (sum(done) / len(done)) * max(0, n - i + 1)
            if remaining >= 5:
                left = f", about {_fmt_secs(remaining)} left"
        total = f" of {n}" if n else ""
        log(f"== Pass {i}{total}: {ctx_text()}{left} ==")


def ctx_pop(tables=None):
    """End the current pass. For each table name in `tables` (Add row blocks of this loop that
    keep skipped passes), a pass that added no row gets one with the loop values and the reason."""
    if not _CTX:
        return
    frame = _CTX[-1]
    for name in tables or ():
        if name not in frame["tables"]:
            tbl_fn = globals().get("table")
            if callable(tbl_fn):
                reason = frame["reason"] or "Not reached"
                tbl_fn(name).add({"Status": reason}, with_ctx=True)
    _PASSES["durations"].setdefault(frame["label"], []).append(time.time() - frame["t0"])
    _CTX.pop()


def ctx_dict():
    """The labels of the loops we are inside, outermost first, as an ordered dict."""
    out = {}
    for frame in _CTX:
        label = frame["label"]
        key, n = label, 2
        while key in out:
            key = f"{label} ({n})"
            n += 1
        out[key] = _ctx_value(frame["value"])
    return out


def ctx_tag():
    """A short file-name tag from the loop values, outermost first."""
    return "_".join(str(_ctx_value(f["value"])) for f in _CTX)


def ctx_mark_row(table_name):
    """Note that a table got a row in every pass we are inside (so no skip row is added)."""
    for f in _CTX:
        f["tables"].add(str(table_name))


def case_skipped(msg):
    """A pass ends early on purpose (a guard such as "Load flow did not converge"): log it with
    the loop values, keep the reason for the summary and for the table's skip row."""
    _PASSES["skipped"] += 1
    where = ctx_text()
    if _CTX:
        _CTX[-1]["reason"] = str(msg)
    _PASSES["reasons"].append((where, str(msg)))
    log(f"Skipped{' (' + where + ')' if where else ''}: {msg}")


def _fmt_secs(s):
    s = int(round(s))
    if s < 90:
        return f"{s} s"
    if s < 5400:
        return f"{round(s / 60)} min"
    return f"{s / 3600:.1f} h"


def ctx_text():
    """'Study case = Base, Dispatch = 300' for messages."""
    return ", ".join(f"{k} = {v}" for k, v in ctx_dict().items())


def _ctx_value(v):
    name = getattr(v, "loc_name", None)
    if name is not None:
        return name
    if isinstance(v, float):
        return float(f"{v:.6g}")
    return v


def log_summary():
    """The end-of-run summary: passes run, skipped (with the first reasons), failed, warnings."""
    p = _PASSES
    failed = globals().get("_FAILED", {}).get("count", 0)
    log(f"Passes run: {p['run']}, skipped: {p['skipped']}, failed: {failed}, "
        f"warnings: {_LOG['warnings']}, errors: {_LOG['errors']}.")
    for where, why in p["reasons"][:20]:
        log(f"  skipped {where or '(no loop)'}: {why}")
    if len(p["reasons"]) > 20:
        log(f"  ... and {len(p['reasons']) - 20} more skipped passes (see run.log).")
    files = globals().get("_RUN", {}).get("files", [])
    if files:
        log(f"Files written: {len(files)} in {run_dir()}")
# @end

# @helper case_failed
# @needs log, ctx
# @imports import traceback
_FAILED = {"count": 0}


def case_failed(exc, trial=None):
    """A loop pass (or a search try) raised: log it with where it happened, keep the reason for
    the skip row, mark the try failed, then carry on with the next pass."""
    _FAILED["count"] += 1
    where = ctx_text()
    if _CTX:
        _CTX[-1]["reason"] = f"ERROR: {type(exc).__name__}: {exc}"[:200]
    if trial is not None:
        trial.failed = True
    log(f"This pass failed{' (' + where + ')' if where else ''}: {type(exc).__name__}: {exc}", "error")
    log(traceback.format_exc().rstrip(), "debug")
# @end

# @helper Trial
class Trial:
    """One try of an Optimise or Find limit block: the values tried and what the steps reported.

    A try may report more than once (for example once per contingency inside the try). Then the
    try passes only if every report passed (one failure fails it, an unknown makes it unknown),
    and its score is the worst one for the search's goal."""

    def __init__(self, params, goal="min"):
        self.params = dict(params)
        self.goal = goal
        self.scores = []
        self.oks = []
        self.failed = False
        self.note = ""
        self.stop = False

    def report(self, score=None, ok=None):
        if score is not None:
            try:
                self.scores.append(float(score))
            except (TypeError, ValueError):
                self.oks.append(None)
        if ok is not None or score is None:
            self.oks.append(_tri(ok))
        return self

    @property
    def score(self):
        if not self.scores:
            return None
        return max(self.scores) if self.goal == "min" else min(self.scores)

    @property
    def ok(self):
        """True, False, or None (unknown: nothing reported, or a report that could not say)."""
        if self.failed:
            return None
        if not self.oks:
            return None if not self.scores else True
        if any(v is False for v in self.oks):
            return False
        if any(v is None for v in self.oks):
            return None
        return True

    def stop_search(self):
        """Stop the loop inside a search: finish this try and end the search."""
        self.stop = True
        return self

    def __repr__(self):
        return f"Trial({self.params}, score={self.score}, ok={self.ok})"


def _tri(v):
    """True / False / None from a bool, a number or a TRUE / FALSE / UNDEFINED style text."""
    if v is None:
        return None
    if isinstance(v, str):
        t = v.strip().upper()
        if t in ("TRUE", "YES", "PASS", "OK", "1"):
            return True
        if t in ("FALSE", "NO", "FAIL", "0"):
            return False
        return None
    if isinstance(v, float) and v != v:
        return None
    return bool(v)
# @end

# @helper vars
# @needs ctx
V = {}
V_WHERE = {}


def var_set(name, value, mode="set"):
    """Remember a value under a name. mode: set, append (to a list), min, max, sum, count.
    For min and max, var_where(name) tells which loop pass set the value."""
    if mode == "append":
        V.setdefault(name, []).append(value)
    elif mode == "count":
        V[name] = V.get(name, 0) + 1
    elif mode in ("min", "max", "sum"):
        if value is None:
            return V.get(name)
        old = V.get(name)
        if old is None or (mode == "min" and value < old) or (mode == "max" and value > old):
            V[name] = value
            if mode != "sum":
                V_WHERE[name] = ctx_text()
        elif mode == "sum":
            V[name] = old + value
    else:
        V[name] = value
        V_WHERE[name] = ctx_text()
    return V.get(name)


def var_get(name, default=None):
    """The value remembered under name, or default when nothing was remembered."""
    return V.get(name, default)


def var_where(name):
    """Where (which loop values) the remembered value was set, as text."""
    return V_WHERE.get(name, "")
# @end

# @helper convert
# @imports import math
def as_float(x, default=None):
    """A number from x, or default when x is empty or not a number."""
    if x is None:
        return default
    if isinstance(x, bool):
        return 1.0 if x else 0.0
    if isinstance(x, (int, float)):
        return float(x)
    if hasattr(x, "item") and not isinstance(x, (list, tuple, str)):
        try:
            return float(x.item())
        except Exception:
            pass
    if isinstance(x, (list, tuple)) and len(x) == 1:
        return as_float(x[0], default)
    try:
        return float(str(x).strip().replace(",", "."))
    except Exception:
        return default


def as_list(x):
    """A list of numbers from a list, a time series (its values), one number, or nothing."""
    if x is None:
        return []
    y = getattr(x, "y", None)
    if y is not None and hasattr(x, "t"):
        return [float(v) for v in y]
    if isinstance(x, (list, tuple)) or (hasattr(x, "__iter__") and not isinstance(x, (str, dict))):
        out = []
        for v in x:
            f = as_float(v)
            if f is not None:
                out.append(f)
        return out
    f = as_float(x)
    return [] if f is None else [f]


def as_items(x):
    """A list of anything (numbers, names, elements) from a list, a series (its values), a
    single value, or nothing. Unlike as_list it keeps text and elements."""
    if x is None:
        return []
    y = getattr(x, "y", None)
    if y is not None and hasattr(x, "t"):
        return list(y)
    if isinstance(x, (list, tuple)):
        return list(x)
    if isinstance(x, (str, bytes, dict)):
        return [x]
    if hasattr(x, "__iter__"):
        try:
            return list(x)
        except Exception:
            return [x]
    return [x]


def as_bool(x):
    """True or False from a value: text such as yes / true / pass / 1 counts as True, NaN as False."""
    if isinstance(x, str):
        return x.strip().lower() in ("1", "true", "yes", "y", "pass", "on")
    if isinstance(x, float) and math.isnan(x):
        return False
    return bool(x)


def as_text(x):
    """Readable text for a value: numbers to 6 significant digits, objects by name."""
    if x is None:
        return ""
    if isinstance(x, bool):
        return "Yes" if x else "No"
    if isinstance(x, float):
        if math.isnan(x):
            return "nan"
        return f"{x:.6g}"
    name = getattr(x, "loc_name", None)
    if name is not None:
        return str(name)
    if isinstance(x, (list, tuple)):
        return ", ".join(as_text(v) for v in x)
    return str(x)
# @end

# @helper num_range
# @needs convert
def num_range(start, stop, step=1):
    """The numbers from start to stop (both included) in steps; rounded to the step's decimals so
    0.1 steps give 0.3, not 0.30000000000000004. Whole numbers stay int."""
    a, b, s = as_float(start), as_float(stop), as_float(step)
    if a is None or b is None:
        return []
    if not s:
        return [a] if a == b else []
    if (b - a) * s < 0:
        return []
    text = repr(float(s))
    decimals = 0
    if "e-" in text:
        decimals = int(text.split("e-")[1]) + len(text.split("e-")[0].split(".")[-1]) if "." in text else int(text.split("e-")[1])
    elif "." in text:
        decimals = len(text.split(".")[1].rstrip("0"))
    for v in (start, step):
        f = as_float(v)
        if f is not None and "." in repr(f):
            decimals = max(decimals, len(repr(f).split(".")[1].rstrip("0")))
    n = int(math.floor((b - a) / s + 1e-9))
    whole = all(isinstance(v, int) and not isinstance(v, bool) for v in (start, stop, step))
    out = []
    for k in range(n + 1):
        v = a + k * s
        out.append(int(round(v)) if whole else round(v, min(decimals, 12)))
    return out
# @end

# @helper compare
# @needs convert
def compare(a, op, b=None, tol=1e-9):
    """True / False, or None when a value is missing. op: "true" (a is true: yes, pass, 1 ...),
    "gt", "ge", "lt", "le", "eq", "ne" (numbers; = and != within tol)."""
    if op == "true":
        if a is None:
            return None
        if isinstance(a, str):
            t = a.strip().upper()
            if t in ("TRUE", "YES", "Y", "PASS", "OK", "ON", "1"):
                return True
            if t in ("FALSE", "NO", "N", "FAIL", "OFF", "0"):
                return False
            return None
        if isinstance(a, float) and a != a:
            return None
        return bool(a)
    x, y = as_float(a), as_float(b)
    if x is None or y is None or x != x or y != y:
        return None
    t = abs(tol or 0.0)
    if op == "gt":
        return x > y
    if op == "ge":
        return x >= y
    if op == "lt":
        return x < y
    if op == "le":
        return x <= y
    if op == "eq":
        return abs(x - y) <= t
    if op == "ne":
        return abs(x - y) > t
    raise ValueError(f"unknown comparison {op!r}")
# @end

# @helper elementwise
# @needs convert
def elementwise(fn, *args):
    """fn applied value by value when any argument is a list (lists of the same length, or lists
    and single numbers); None where a value is missing or the result is not a number."""
    lists = [a for a in args if isinstance(a, (list, tuple)) or (hasattr(a, "y") and hasattr(a, "t"))]

    def one(*xs):
        if any(x is None for x in xs):
            return None
        try:
            return fn(*xs)
        except (ArithmeticError, ValueError, TypeError):
            return None
    if not lists:
        return one(*[as_float(a) for a in args])
    seqs = [as_list(a) if (isinstance(a, (list, tuple)) or hasattr(a, "y")) else None for a in args]
    n = min(len(s) for s in seqs if s is not None)
    out = []
    for k in range(n):
        out.append(one(*[(s[k] if s is not None else as_float(a)) for s, a in zip(seqs, args)]))
    return out
# @end

# @helper arith
# @needs elementwise
def arith(op, a, b):
    """a op b for numbers (or value by value for lists): add, sub, mul, div, pow, min, max, mod."""
    fns = {
        "add": lambda x, y: x + y, "sub": lambda x, y: x - y, "mul": lambda x, y: x * y,
        "div": lambda x, y: x / y, "pow": lambda x, y: x ** y, "min": lambda x, y: min(x, y),
        "max": lambda x, y: max(x, y), "mod": lambda x, y: x % y,
    }
    return elementwise(fns[op], a, b)
# @end

# @helper func1
# @needs elementwise
def func1(fn, x, digits=None):
    """One function of a number (or of each value of a list): abs, round, sqrt, neg, floor,
    ceil, log10, exp."""
    fns = {
        "abs": abs, "round": lambda v: round(v, int(digits or 0)), "sqrt": math.sqrt,
        "neg": lambda v: -v, "floor": math.floor, "ceil": math.ceil, "log10": math.log10, "exp": math.exp,
    }
    return elementwise(fns[fn], x)
# @end

# @helper logic
# @needs convert
def logic(op, a, b=None):
    """and / or / not on yes-no values, with None for unknown: and is False when either is
    False, or is True when either is True; otherwise an unknown makes the answer unknown."""
    def tri(v):
        if v is None:
            return None
        if isinstance(v, str):
            t = v.strip().upper()
            return True if t in ("TRUE", "YES", "PASS", "OK", "1") else False if t in ("FALSE", "NO", "FAIL", "0") else None
        if isinstance(v, float) and v != v:
            return None
        return bool(v)
    x, y = tri(a), tri(b)
    if op == "not":
        return None if x is None else not x
    if op == "and":
        if x is False or y is False:
            return False
        return None if x is None or y is None else True
    if op == "or":
        if x is True or y is True:
            return True
        return None if x is None or y is None else False
    raise ValueError(f"unknown logic {op!r}")
# @end

# @helper apply_formula
# @needs elementwise, log
def apply_formula(fn, args, each=False, label="Formula"):
    """A formula block's lambda applied to its inputs: None when an input is missing or the
    maths fails (with a warning), value by value for lists when each is True."""
    if each:
        return elementwise(fn, *args)
    if any(a is None for a in args):
        return None
    try:
        vals = [as_float(a, a) for a in args]
        return fn(*vals)
    except (ArithmeticError, ValueError, TypeError) as exc:
        log(f"{label}: could not work it out ({exc}); the result is empty.", "warn")
        return None
# @end

# @helper text_format
# @needs convert
# @imports import string
def text_format(template, values):
    """The template with each {name} replaced by its value: numbers to 6 significant digits,
    elements by name, {name:.2f} with that format; an empty value gives an empty piece."""
    try:
        parts = list(string.Formatter().parse(str(template)))
    except ValueError:
        return str(template)
    out = []
    for text, field, spec, conv in parts:
        out.append(text)
        if field is None:
            continue
        v = values.get(field)
        if spec and v is not None:
            try:
                f = as_float(v) if not isinstance(v, str) else v
                out.append(format(f if f is not None else v, spec))
                continue
            except (ValueError, TypeError):
                pass
        out.append(as_text(v))
    return "".join(out)
# @end

# @helper Series
# @imports import math
class Series:
    """A time series: t in seconds and y values, same length, from a simulation result.
    `events` lists the workflow's simulation events of that run as (time in s, label), so the
    checks can find the event time and the graphs can mark it."""

    __slots__ = ("t", "y", "label", "unit", "events")

    def __init__(self, t, y, label="", unit="", events=None):
        t = [float(v) for v in t]
        y = [float(v) if v is not None else float("nan") for v in y]
        n = min(len(t), len(y))
        self.t, self.y, self.label, self.unit = t[:n], y[:n], str(label), str(unit)
        self.events = [(float(et), str(el)) for et, el in (events or [])]

    def __len__(self):
        return len(self.t)

    def __bool__(self):
        return len(self.t) > 0

    def first_event(self, default=None):
        """Time of the first workflow event of the run (s), or default."""
        return min((et for et, _ in self.events), default=default)

    def between(self, t0=None, t1=None):
        """The part with t0 <= t <= t1 (either end may be None)."""
        keep = [i for i, tv in enumerate(self.t)
                if (t0 is None or tv >= t0) and (t1 is None or tv <= t1)]
        return Series([self.t[i] for i in keep], [self.y[i] for i in keep], self.label, self.unit, self.events)

    def after(self, t0):
        return self.between(t0, None)

    def finite(self):
        """The samples whose value is a real number (NaN and inf dropped)."""
        keep = [i for i, v in enumerate(self.y) if math.isfinite(v)]
        return Series([self.t[i] for i in keep], [self.y[i] for i in keep], self.label, self.unit, self.events)

    def __repr__(self):
        span = f"{self.t[0]:g}..{self.t[-1]:g} s" if self.t else "empty"
        return f"Series({self.label or 'unnamed'}, {len(self)} samples, {span})"
# @end

# @helper table
# @needs log, run_dir, ctx, convert
# @imports import csv
_TABLES = {}


class ResultTable:
    """Rows of results. Columns appear in the order they are first used."""

    def __init__(self, name):
        self.name = str(name)
        self.rows = []
        self.columns = []
        self.saved_rows = 0
        self.paths = []

    def add(self, row, with_ctx=True, status=False):
        """Add a row: the loop values first (with_ctx), then the row's own columns. With status,
        a Status column says OK (skip rows written by ctx_pop carry the reason instead)."""
        full = {}
        if with_ctx:
            full.update(ctx_dict())
        if status:
            full["Status"] = "OK"
        full.update(row)
        for k in full:
            if k not in self.columns:
                self.columns.append(k)
        self.rows.append({k: _cell(v) for k, v in full.items()})
        ctx_mark_row(self.name)
        return self

    def column(self, name):
        """The values of one column, top to bottom (empty cells left out)."""
        return [r[name] for r in self.rows if r.get(name) not in (None, "")]

    def save(self, stem=None, fmt="csv"):
        """Write the table to the run folder as csv, xlsx or both; returns the paths."""
        stem = stem or self.name
        paths = []
        if fmt in ("csv", "both"):
            path = out_file(stem, "csv")
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(self.columns)
                for r in self.rows:
                    w.writerow(["" if r.get(c) is None else r.get(c) for c in self.columns])
            paths.append(path)
        if fmt in ("xlsx", "both"):
            try:
                import pandas as pd
                path = out_file(stem, "xlsx")
                pd.DataFrame(self.rows, columns=self.columns).to_excel(path, index=False)
                paths.append(path)
            except Exception as e:
                log(f"Could not write an Excel file for '{self.name}' ({e}); writing CSV instead.", "warn")
                if fmt == "xlsx":
                    paths += self.save(stem, "csv")
        for p in paths:
            log(f"Saved table '{self.name}' ({len(self.rows)} rows): {p}")
        self.saved_rows = len(self.rows)
        self.paths += paths
        return paths

    def __len__(self):
        return len(self.rows)


def table(name="Results"):
    """The results table with this name (made on first use)."""
    key = str(name)
    if key not in _TABLES:
        _TABLES[key] = ResultTable(key)
    return _TABLES[key]


def flush_tables():
    """Save every table that has rows no block has saved yet (called when the workflow ends)."""
    for t in _TABLES.values():
        if len(t.rows) > t.saved_rows:
            log(f"Table '{t.name}' has rows that were not saved; saving it now.")
            t.save()


def _cell(v):
    if isinstance(v, bool) or v is None or isinstance(v, (int, str)):
        return v
    if isinstance(v, float):
        return v
    f = as_float(v)
    if f is not None and not isinstance(v, (list, tuple)):
        return f
    return as_text(v)
# @end

# @helper table_pick
# @needs table, convert
def table_pick(name, column, pick="all", of=None, kind="num"):
    """One column of the results table `name`: pick "all" gives its values top to bottom (an
    empty cell as None, so columns of one table stay in step); "max" / "min" give the value of
    `column` in the row where column `of` (default: `column` itself) is highest / lowest, for
    example the contingency with the highest loading. kind "num" gives numbers, "text" texts.
    A table or column that does not exist gives [] (or None)."""
    tbl = _TABLES.get(str(name))
    rows = tbl.rows if tbl is not None else []

    def conv(v):
        if kind == "text":
            return None if v is None or v == "" else as_text(v)
        return as_float(v)

    if pick not in ("max", "min"):
        return [conv(r.get(column)) for r in rows]
    key = of or column
    best = None
    for r in rows:
        v = as_float(r.get(key))
        if v is None or v != v:
            continue
        if best is None or (v > best[0] if pick == "max" else v < best[0]):
            best = (v, r)
    return conv(best[1].get(column)) if best is not None else None
# @end

# @helper read_table
# @needs table, log, convert, StopWorkflow
# @imports import csv
# @imports import io
# @imports import os
# @imports import re
# @imports import sys
def read_table(name, path, sheet="", layout="rows", header=1, decimal_comma=False):
    """Fill the results table `name` from a CSV file or an Excel sheet (.xlsx, .xlsm).

    layout "rows": one case per row, the header row names the columns; "columns": one case per
    column, the first column names the fields. header: the row (for "columns": the column) that
    holds the names, 1 for the first. Numbers become numbers (decimal_comma: 1,5 is 1.5), empty
    cells None, other text stays text (a code such as 007 stays text). Rows with no value are
    left out. The table is filled afresh and counts as saved, so the end of the run does not
    write it to the run folder again. A name without a folder is looked for next to the script
    first (a study package is unpacked as one folder), then in the current folder. Returns the
    table; stops the workflow with a plain message when the file cannot be read."""
    where = _rt_find(path)
    ext = os.path.splitext(where)[1].lower()
    if ext == ".xls":
        raise StopWorkflow(f"Read table from file: {where} is an old .xls file. Save it as .xlsx or .csv in Excel and use that file.")
    try:
        grid = _rt_xlsx(where, sheet) if ext in (".xlsx", ".xlsm") else _rt_csv(where)
    except StopWorkflow:
        raise
    except Exception as exc:
        raise StopWorkflow(f"Read table from file: could not read {where}: {type(exc).__name__}: {exc}")
    if str(layout or "rows").strip().lower() == "columns":
        width = max((len(r) for r in grid), default=0)
        grid = [[r[k] if k < len(r) else None for r in grid] for k in range(width)]
    h = max(1, int(as_float(header, 1) or 1))
    if len(grid) < h:
        raise StopWorkflow(f"Read table from file: {where} has no {'column' if layout == 'columns' else 'row'} {h} with the names.")
    names, seen = [], {}
    for k, v in enumerate(grid[h - 1]):
        nm = "" if v is None else str(v).strip()
        nm = nm or f"Column {k + 1}"
        c = seen.get(nm, 0)
        seen[nm] = c + 1
        names.append(nm if not c else f"{nm} ({c + 1})")
    rows = []
    for raw in grid[h:]:
        vals = [_rt_value(v, decimal_comma) for v in raw]
        if all(v is None for v in vals):
            continue
        if len(vals) > len(names):
            for k in range(len(names), len(vals)):
                names.append(f"Column {k + 1}")
        rows.append({names[k]: (vals[k] if k < len(vals) else None) for k in range(len(names))})
    # Columns with no name and no value (trailing blanks of a sheet) are left out.
    keep = [nm for k, nm in enumerate(names) if not nm.startswith("Column ") or any(r.get(nm) is not None for r in rows)]
    tbl = table(name)
    tbl.rows, tbl.columns = [], []
    for r in rows:
        tbl.add({nm: r.get(nm) for nm in keep}, with_ctx=False)
    tbl.columns = list(keep) if not rows else tbl.columns
    tbl.saved_rows = len(tbl.rows)
    log(f"Read {len(tbl.rows)} rows with {len(keep)} columns from {where} into table '{tbl.name}'.")
    return tbl


def _rt_find(path):
    p = os.path.expanduser(str(path or "").strip().strip('"'))
    if not p:
        raise StopWorkflow("Read table from file: no file is given. Fill in the file on the block (or its line in SETTINGS).")
    tried = []
    if os.path.isabs(p):
        tried.append(p)
    else:
        # Next to the script first. A study package is unpacked as one folder and may be run from
        # anywhere, so the file that travelled with the script must win over one that happens to
        # have the same name in the folder the run was started from; that folder is still tried.
        here = globals().get("__file__") or (sys.argv[0] if sys.argv and sys.argv[0] else "")
        if here:
            tried.append(os.path.join(os.path.dirname(os.path.abspath(here)), p))
        tried.append(os.path.abspath(p))
    for t in tried:
        if os.path.isfile(t):
            return t
    raise StopWorkflow("Read table from file: cannot find " + " or ".join(dict.fromkeys(tried)) + ".")


def _rt_csv(path):
    with open(path, "rb") as fh:
        data = fh.read()
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    sample = text[:4096]
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
    except csv.Error:
        first = sample.splitlines()[0] if sample else ""
        delim = max((";", ",", "\t"), key=first.count)
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]


def _rt_xlsx(path, sheet=""):
    try:
        import openpyxl
    except ImportError:
        raise StopWorkflow("Read table from file: reading an Excel file needs the openpyxl package "
                           "(pip install openpyxl), or save the sheet as CSV.")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        want = str(sheet or "").strip()
        if want and want not in wb.sheetnames:
            raise StopWorkflow(f"Read table from file: {path} has no sheet '{want}' (it has: {', '.join(wb.sheetnames)}).")
        ws = wb[want] if want else wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _rt_value(v, decimal_comma=False):
    # A cell as a number, None (empty) or text.
    if v is None or isinstance(v, (bool, int, float)):
        return v
    if hasattr(v, "isoformat"):
        return v.isoformat()
    s = str(v).strip()
    if not s:
        return None
    t = s.replace(" ", "").replace(" ", "")
    if decimal_comma:
        t = t.replace(".", "").replace(",", ".")
    if re.fullmatch(r"[+-]?0\d+", t):
        return s
    if re.fullmatch(r"[+-]?\d+", t):
        return int(t)
    if re.fullmatch(r"[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?", t):
        return float(t)
    return s
# @end
